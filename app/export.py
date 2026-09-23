"""CSV and Excel export endpoints.

Implements the P1 items in OUTPUT-LAYER-PLAN.md:
    GET /export/documents.csv
    GET /export/documents.xlsx
    GET /export/ledger.csv
    GET /export/ledger.xlsx

CSV responses carry a UTF-8 BOM so Excel auto-detects the encoding and
non-ASCII entity names render correctly. Content-Disposition triggers a
download rather than an inline render.
"""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.auth import require_user as require_api_key
from app.db import get_session
from app.models.orm import Document, Instrument, LedgerEntry


router = APIRouter(prefix="/export", tags=["export"])


def _enum_val(x: Any) -> str:
    if x is None:
        return ""
    return x.value if hasattr(x, "value") else str(x)


def _csv_bytes(headers: list, rows: list) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return buf.getvalue().encode("utf-8-sig")


def _xlsx_bytes(headers: list, rows: list, sheet: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([None if v == "" else v for v in row])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(h) + 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_response(filename: str, body: bytes) -> Response:
    return Response(
        content=body,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": "attachment; filename=\"" + filename + "\"",
            "Cache-Control": "no-store",
        },
    )


def _xlsx_response(filename: str, body: bytes) -> Response:
    return Response(
        content=body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=\"" + filename + "\"",
            "Cache-Control": "no-store",
        },
    )


def _document_rows(session, issuer_name):
    q = session.query(Document)
    if issuer_name:
        q = q.join(Instrument, Document.instrument_id == Instrument.id).filter(
            Instrument.issuer_name == issuer_name
        )
    docs = q.all()
    headers = [
        "document_id", "filename", "instrument_id", "document_type",
        "classification_confidence", "extraction_confidence",
        "extraction_schema_name", "status", "uploaded_at",
    ]
    rows = []
    for d in docs:
        rows.append([
            getattr(d, "id", ""),
            getattr(d, "filename", ""),
            getattr(d, "instrument_id", ""),
            _enum_val(getattr(d, "document_type", None)),
            getattr(d, "classification_confidence", ""),
            getattr(d, "extraction_confidence", ""),
            getattr(d, "extraction_schema_name", "") or "",
            _enum_val(getattr(d, "status", None)),
            str(getattr(d, "uploaded_at", "") or ""),
        ])
    return headers, rows


def _ledger_rows(session, instrument_id, issuer_name):
    q = session.query(LedgerEntry)
    if instrument_id:
        q = q.filter(LedgerEntry.instrument_id == instrument_id)
    elif issuer_name:
        inst_ids = [
            row[0]
            for row in session.query(Instrument.id)
            .filter(Instrument.issuer_name == issuer_name)
            .all()
        ]
        if inst_ids:
            q = q.filter(LedgerEntry.instrument_id.in_(inst_ids))
        else:
            q = q.filter(LedgerEntry.instrument_id == "__none__")
    entries = q.all()
    headers = [
        "entry_id", "entry_type", "instrument_id", "document_id",
        "event", "trace_id", "created_at", "payload_json",
    ]
    rows = []
    for e in entries:
        payload = getattr(e, "payload", None) or {}
        if not isinstance(payload, dict):
            payload = {}
        ts = getattr(e, "created_at", None) or getattr(e, "timestamp", None) or ""
        rows.append([
            getattr(e, "id", ""),
            _enum_val(getattr(e, "entry_type", None)),
            getattr(e, "instrument_id", "") or "",
            getattr(e, "document_id", "") or "",
            payload.get("event", ""),
            payload.get("trace_id", ""),
            str(ts),
            json.dumps(payload, default=str),
        ])
    return headers, rows


@router.get("/documents.csv")
def export_documents_csv(
    issuer_name: str | None = Query(default=None),
    session: Session = Depends(get_session),
    _: Any = Depends(require_api_key),
) -> Response:
    h, r = _document_rows(session, issuer_name)
    return _csv_response("flowgate-documents.csv", _csv_bytes(h, r))


@router.get("/documents.xlsx")
def export_documents_xlsx(
    issuer_name: str | None = Query(default=None),
    session: Session = Depends(get_session),
    _: Any = Depends(require_api_key),
) -> Response:
    h, r = _document_rows(session, issuer_name)
    return _xlsx_response("flowgate-documents.xlsx", _xlsx_bytes(h, r, "Documents"))


@router.get("/ledger.csv")
def export_ledger_csv(
    instrument_id: str | None = Query(default=None),
    issuer_name: str | None = Query(default=None),
    session: Session = Depends(get_session),
    _: Any = Depends(require_api_key),
) -> Response:
    h, r = _ledger_rows(session, instrument_id, issuer_name)
    return _csv_response("flowgate-ledger.csv", _csv_bytes(h, r))


@router.get("/ledger.xlsx")
def export_ledger_xlsx(
    instrument_id: str | None = Query(default=None),
    issuer_name: str | None = Query(default=None),
    session: Session = Depends(get_session),
    _: Any = Depends(require_api_key),
) -> Response:
    h, r = _ledger_rows(session, instrument_id, issuer_name)
    return _xlsx_response("flowgate-ledger.xlsx", _xlsx_bytes(h, r, "Ledger"))
